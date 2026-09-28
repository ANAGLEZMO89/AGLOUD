"""Prueba de la extensión contra la MAQUETA local (datos ficticios, no es Wtransnet).

Uso: python pruebas/probar_extension.py
Carga la extensión en Chromium, abre la maqueta con «sesión», rellena el panel,
pulsa Prueba/Buscar/Revisar y comprueba el Excel descargado y que no se ha
pedido ninguna URL comercial.
"""
import json, os, shutil, sys, tempfile, time
from urllib.parse import parse_qs, urlsplit

AQUI = os.path.dirname(os.path.abspath(__file__))
EXT = os.path.dirname(AQUI)
sys.path.insert(0, os.path.join(os.path.dirname(EXT), "wtransnet_bot"))
from tests.maqueta import servidor  # noqa: E402
import openpyxl  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402


def preparar_copia():
    d = tempfile.mkdtemp()
    dst = os.path.join(d, "ext")
    shutil.copytree(EXT, dst, ignore=shutil.ignore_patterns("pruebas"))
    m = json.load(open(os.path.join(dst, "manifest.json")))
    m["host_permissions"].append("http://127.0.0.1/*")  # sólo en la copia de prueba
    json.dump(m, open(os.path.join(dst, "manifest.json"), "w"))
    return dst


def main(js=False, real=False):
    srv = servidor.arrancar()
    servidor.MODO["enlaces_js"] = js
    servidor.MODO["estilo_real"] = real
    base = f"http://127.0.0.1:{srv.server_port}/WTNWEB/"
    ext = preparar_copia()
    salida = tempfile.mkdtemp()
    fallos = []
    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(tempfile.mkdtemp(), executable_path="/opt/pw-browsers/chromium-1194/chrome-linux/chrome", headless=True, accept_downloads=True,
                                                   args=[f"--disable-extensions-except={ext}", f"--load-extension={ext}"])
        sw = ctx.service_workers[0] if ctx.service_workers else ctx.wait_for_event("serviceworker")
        ext_id = sw.url.split("/")[2]
        wt = ctx.pages[0] if ctx.pages else ctx.new_page()
        wt.goto(base)
        panel = ctx.new_page()
        errores_js = []
        panel.on("pageerror", lambda e: errores_js.append(str(e)))
        panel.goto(f"chrome-extension://{ext_id}/panel.html")
        panel.evaluate(f"chrome.storage.local.set({{baseUrl: '{base}'}})")
        panel.fill("[name=origenes]", "España; Valencia")
        if real:
            panel.select_option("[name=ambito_origen]", "Provincia y colindantes")
        panel.fill("[name=fecha_inicial]", "28/09/26")
        panel.fill("[name=fecha_final]", "05/10/26")
        panel.check("input[name=tipo_bolsa][value='Trailers Completos']")
        panel.check("input[name=redes][value=Wtransnet]")
        panel.check("input[name=redes][value=Teleroute]")
        panel.select_option("[name=adr]", "Indiferente")
        panel.fill("[name=empresa_esperada]", "EMPRESA PRUEBA MAQUETA")
        panel.click("details summary")
        panel.fill("[name=pausa_min]", "2")
        panel.fill("[name=pausa_max]", "3")
        servidor.PETICIONES.clear()

        # 1) Revisión de la página
        panel.click("#btnInspeccionar")
        fin = time.time() + 180
        while time.time() < fin and not panel.evaluate("() => !document.querySelector('#zonaInforme').hidden || document.querySelector('#estado').className.includes('error')"):
            time.sleep(0.5)
        informe = panel.input_value("#informe")
        print("ESTADO inspección:", panel.inner_text("#estado"))
        print("\n".join(l for l in informe.splitlines() if not l.startswith("CONTROLES"))[:2500])

        # 2) Ejecución completa
        with panel.expect_download(timeout=600000) as dl:
            panel.click("#btnEjecutar")
        ruta = os.path.join(salida, dl.value.suggested_filename)
        dl.value.save_as(ruta)
        print("ESTADO ejecución:", panel.inner_text("#estado"))
        log = panel.inner_text("#log")
        print(log[-3500:])
        if errores_js:
            fallos.append("Errores JS: " + "; ".join(errores_js))
        ctx.close()

    rutas = [q for _, q in servidor.PETICIONES]
    prohibidas = [q for q in rutas if any(x in q for x in ("accion=alta", "accion=interes", "/chat", "mailto", "tel:"))]
    if prohibidas:
        fallos.append(f"Peticiones prohibidas: {prohibidas}")
    listar = [parse_qs(urlsplit(q).query) for q in rutas if "accion=listar" in q and "cgcm=CG" in q and ("fIni" in q or "FechaDisp" in q)]
    if real:
        ok_f = listar and listar[0].get("province_from") == ["5"] and listar[0].get("FechaDisp") == ["28/09/26"] and listar[0].get("TipoBolsa_Completa") == ["1"] \
            and listar[0].get("region_from") == ["1"]
    else:
        ok_f = listar and listar[0].get("provOri") == ["3"] and listar[0].get("fIni") == ["28/09/26"] and listar[0].get("bTC") == ["1"]
    if not ok_f:
        fallos.append(f"Filtros de carga no enviados correctamente: {listar[:1]}")
    wb = openpyxl.load_workbook(ruta)
    print("HOJAS:", wb.sheetnames)
    cargas = list(wb["Cargas"].iter_rows(values_only=True))
    ids = [r[0] for r in cargas[1:]]
    print("CARGAS:", ids)
    if ids != ["Wtransnet#100001", "Teleroute#200002"]:
        fallos.append(f"Cargas inesperadas: {ids}")
    cab = cargas[0]
    c1 = dict(zip(cab, cargas[1])) if len(cargas) > 1 else {}
    for k, esperado in [("Teléfonos", "+34961000001"), ("Emails", "trafico@maqueta-uno.test"), ("Precio (texto)", "0 EUR"), ("Origen CP", "46001"), ("Peso kg", 24000)]:
        if c1.get(k) != esperado:
            fallos.append(f"{k}: {c1.get(k)!r} != {esperado!r}")
    if "No contestamos correos" not in str(c1.get("Advertencias de contacto")):
        fallos.append("Falta aviso de contacto")
    m = list(wb["Matching"].iter_rows(values_only=True))
    rel = [dict(zip(m[0], r)) for r in m[1:]]
    print("MATCHING:", [(r["ID carga"], r["ID camión"], r["Estado general"], r["Cobertura (interna)"]) for r in rel])
    if not any(r["ID camión"] == "Wtransnet#500001" and r["Capacidad y dimensiones"] == "Compatible según datos publicados" for r in rel):
        fallos.append("Relación K1 no encontrada o capacidad mal evaluada")
    if any(r["ID camión"] == "Wtransnet#500002" for r in rel):
        fallos.append("El furgón de 1 Tn no debería estar en Matching")
    inc = " ".join(str(r[3]) for r in wb["Incidencias"].iter_rows(min_row=2, values_only=True))
    for t in ("No vigente", "1000 kg", "posible misma oferta"):
        if t not in inc:
            fallos.append(f"Incidencias sin «{t}»")
    c2 = dict(zip(cab, cargas[2])) if len(cargas) > 2 else {}
    if not str(c2.get("Fuente del contacto", "")).startswith("Contacto GENERAL"):
        fallos.append("Contacto general de empresa no marcado en C2")
    ws = wb["Matching"]
    if ws.freeze_panes is None or not ws.auto_filter.ref:
        fallos.append("Matching sin inmovilizar o sin filtro")
    srv.shutdown()
    print("\nRESULTADO:", "OK" if not fallos else "FALLOS:\n- " + "\n- ".join(fallos))
    return not fallos


if __name__ == "__main__":
    ok = main(js="--js" in sys.argv, real="--real" in sys.argv)
    sys.exit(0 if ok else 1)
