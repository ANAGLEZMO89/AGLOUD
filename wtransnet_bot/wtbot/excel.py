"""Exportación a Excel (.xlsx) con openpyxl.

- Encabezados inmovilizados, autofiltro, ajuste de texto y anchos legibles.
- Fechas como fecha real de Excel; CP y teléfonos como texto.
- Protección frente a fórmulas: todo texto procedente de terceros se escribe
  como texto (tipo cadena + prefijo de comilla), aunque empiece por = + - @.
- Enlaces a fichas sin parámetros de sesión.
"""
from __future__ import annotations

import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import compatibilidad as C
from .duplicados import posibles_duplicados
from .navegador import url_limpia

CAB = PatternFill("solid", fgColor="1F4E78")
FUENTE_CAB = Font(bold=True, color="FFFFFF")
COLOR_ESTADO = {C.COMPATIBLE: "C6EFCE", C.PENDIENTE: "FFEB9C", C.INCOMPATIBLE: "FFC7CE", "Descartado": "FFC7CE"}


def _texto(ws, fila, col, valor):
    c = ws.cell(row=fila, column=col)
    if valor is None or valor == "":
        c.value = None
        return c
    if isinstance(valor, bool):
        valor = "Sí" if valor else "No"
    if isinstance(valor, dt.datetime):
        c.value = valor
        c.number_format = "dd/mm/yyyy hh:mm" if (valor.hour or valor.minute) else "dd/mm/yyyy"
        return c
    if isinstance(valor, (int, float)):
        c.value = valor
        return c
    if isinstance(valor, (list, tuple)):
        valor = "\n".join(str(v) for v in valor if v not in (None, ""))
    s = ILLEGAL_CHARACTERS_RE.sub("", str(valor))[:32000]
    c.value = s
    c.data_type = "s"  # nunca fórmula
    if s[:1] in ("=", "+", "-", "@", "\t", "\r"):
        c.quotePrefix = True
    return c


def _enlace(ws, fila, col, url, texto="Abrir ficha"):
    url = url_limpia(url or "")
    if not url:
        return
    c = ws.cell(row=fila, column=col, value=texto)
    c.hyperlink = url
    c.font = Font(color="0563C1", underline="single")


def _hoja(wb, titulo, columnas, anchos=None):
    ws = wb.create_sheet(titulo)
    for i, col in enumerate(columnas, 1):
        c = ws.cell(row=1, column=i, value=col)
        c.fill, c.font = CAB, FUENTE_CAB
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = (anchos or {}).get(col, max(12, min(40, len(col) + 4)))
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 32
    return ws


def _cerrar_hoja(ws):
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            if c.hyperlink is None:
                c.alignment = Alignment(wrap_text=True, vertical="top")


def _ubic(u):
    return u.get("texto", "") if u else ""


def _cant(q):
    if not q or not q.get("original"):
        return ""
    return q["original"]


def _num(q):
    return q.get("valor") if q else None


def _sn(v):
    return {True: "Sí", False: "No", None: "No consta"}[v]


COLS_OFERTA = [
    "ID bot (red#nº)", "Red", "Nº oferta", "Enlace ficha", "Leído el", "Modificada", "Disponibilidad (texto)",
    "Disponible desde", "Disponible hasta", "Hora límite", "Fecha descarga", "Origen (texto)", "Origen país",
    "Origen provincia", "Origen CP", "Origen localidad", "Origen ámbito", "Destino (texto)", "Destino país",
    "Destino provincia", "Destino CP", "Destino localidad", "Destino ámbito", "Tipo de bolsa", "Vehículo",
    "Especialidad", "Peso (texto)", "Peso kg", "Volumen (texto)", "Volumen m3", "Largo (texto)", "Ancho (texto)",
    "Alto (texto)", "Forma de carga", "Mercancía / características", "ADR", "Plataforma elevadora",
    "Doble conductor", "Otros equipamientos", "Observaciones (texto original)", "Nº viajes", "Ida y vuelta",
    "Distancia (si consta)", "Precio (texto)", "Precio valor", "Moneda", "Plazo de pago", "Forma de pago",
    "Comentarios de pago", "Código empresa", "Empresa", "Persona de contacto", "Teléfonos", "Móviles", "Emails",
    "Fuente del contacto", "Advertencias de contacto", "Actividad empresarial", "Advertencias", "Otros datos de la ficha",
]
ANCHOS = {"Observaciones (texto original)": 60, "Otros datos de la ficha": 60, "Advertencias": 45,
          "Advertencias de contacto": 45, "Origen (texto)": 30, "Destino (texto)": 30, "Empresa": 30,
          "Mercancía / características": 35, "Motivo de la coincidencia": 70, "Información pendiente": 70,
          "Detalle": 90, "Valor": 90}


def _fila_oferta(r: dict) -> list:
    c = r["campos"]
    o, d = r.get("origen", {}), r.get("destino", {})
    return [
        r["id"], r.get("red") or "No indicada", r.get("numero_oferta"), None, r.get("leido_en"),
        r.get("fecha_modificacion") or c.get("fecha_modificacion"), c.get("disponibilidad"), r.get("disp_desde"),
        r.get("disp_hasta"), r.get("hora_limite") or c.get("hora_limite"), r.get("fecha_descarga") or c.get("fecha_descarga"),
        _ubic(o), o.get("pais") or o.get("pais_iso"), o.get("provincia") or o.get("provincia_cod"), o.get("cp"),
        o.get("localidad"), o.get("ambito"), _ubic(d), d.get("pais") or d.get("pais_iso"),
        d.get("provincia") or d.get("provincia_cod"), d.get("cp"), d.get("localidad"), d.get("ambito"),
        c.get("tipo_bolsa"), c.get("vehiculo"), c.get("especialidad"), _cant(r.get("peso")), _num(r.get("peso")),
        _cant(r.get("volumen")), _num(r.get("volumen")), _cant(r.get("largo")), _cant(r.get("ancho")),
        _cant(r.get("alto")), c.get("forma_carga"), c.get("mercancia"), c.get("adr") or _sn(r.get("adr")),
        c.get("plataforma_elevadora") or _sn(r.get("plataforma_elevadora")),
        c.get("doble_conductor") or _sn(r.get("doble_conductor")), c.get("equipamiento"), c.get("observaciones"),
        c.get("num_viajes"), c.get("ida_y_vuelta"), c.get("distancia"), r.get("precio", {}).get("original"),
        r.get("precio", {}).get("valor"), r.get("precio", {}).get("moneda") or c.get("moneda"), c.get("plazo_pago"),
        c.get("forma_pago"), c.get("comentarios_pago"), r.get("empresa_codigo"), r.get("empresa_nombre"),
        r.get("contacto"), r.get("telefonos"), r.get("moviles"), r.get("emails"), _fuente(r),
        r.get("avisos_contacto"), c.get("actividad") or r.get("empresa_ficha", {}).get("actividad"),
        r.get("avisos"), r.get("otros_datos"),
    ]


def _fuente(r):
    f = r.get("contacto_fuente")
    if f == "oferta":
        return "Contacto de la oferta"
    if f == "empresa":
        return "Contacto GENERAL de la empresa (ficha de empresa), no de la oferta"
    return "Sin contacto visible"


def _escribir_ofertas(ws, registros):
    ienlace = COLS_OFERTA.index("Enlace ficha") + 1
    for i, r in enumerate(registros, 2):
        for j, v in enumerate(_fila_oferta(r), 1):
            if j != ienlace:
                _texto(ws, i, j, v)
        _enlace(ws, i, ienlace, r.get("url"))


def exportar(progreso: dict, carpeta: str, ahora: dt.datetime | None = None) -> str:
    ahora = ahora or dt.datetime.now()
    os.makedirs(carpeta, exist_ok=True)
    ruta = os.path.join(carpeta, f"Wtransnet_cargas_camiones_{ahora:%Y%m%d_%H%M}.xlsx")
    cargas = [progreso["cargas"][i] for i in progreso["orden_cargas"] if i in progreso["cargas"]]
    camiones_usados = []
    vistos = set()
    for cid in progreso["orden_cargas"]:
        for rel in progreso["relaciones"].get(cid, []):
            if rel["camion"] not in vistos and rel["camion"] in progreso["camiones"]:
                vistos.add(rel["camion"])
                camiones_usados.append(progreso["camiones"][rel["camion"]])
    wb = Workbook()
    wb.remove(wb.active)

    # ---------------- Matching (primero en importancia, pero Resumen va delante)
    rels_ok = [(cid, rel) for cid in progreso["orden_cargas"] for rel in progreso["relaciones"].get(cid, [])
               if rel["estado"] != "Descartado"]
    n_comp = sum(1 for _, r in rels_ok if r["estado"] == C.COMPATIBLE)
    n_pend = sum(1 for _, r in rels_ok if r["estado"] == C.PENDIENTE)
    n_desc = sum(1 for cid in progreso["orden_cargas"] for r in progreso["relaciones"].get(cid, []) if r["estado"] == "Descartado")
    empresas = {(t.get("empresa_codigo") or t.get("empresa_nombre") or "").strip().lower() for t in camiones_usados} - {""}

    ws = _hoja(wb, "Resumen", ["Concepto", "Valor"], {"Concepto": 42, "Valor": 90})
    filas = [
        ("Fecha y hora de ejecución", progreso.get("inicio")),
        ("Fecha y hora de exportación", ahora),
        ("Estado de la ejecución", progreso.get("estado")),
        ("Cuenta visible en la sesión", progreso.get("empresa_sesion") or "No verificada"),
        ("Criterios solicitados (filtros Wtransnet)", _criterios_txt(progreso["criterios"].get("filtros_wtransnet", {}))),
        ("Comprobaciones posteriores del bot", _criterios_txt(progreso["criterios"].get("comprobaciones_bot", {}))),
        ("Límites", _criterios_txt(progreso["criterios"].get("limites", {}))),
        ("Filtros aplicados y verificados en el formulario de cargas", progreso.get("filtros_aplicados", {}).get("carga")),
        ("Cargas incluidas", len(cargas)),
        ("Cargas leídas y excluidas por comprobaciones", len(progreso.get("descartadas_carga", []))),
        ("Ofertas de camión en relaciones", len(camiones_usados)),
        ("Empresas distintas de camión", len(empresas)),
        ("Vehículos identificados", "No calculado: las fichas no identifican vehículos de forma inequívoca (matrícula)"),
        ("Relaciones compatibles según datos publicados", n_comp),
        ("Relaciones pendientes de confirmar", n_pend),
        ("Candidatos descartados por incompatibilidad", f"{n_desc} (ver hoja Incidencias)"),
        ("Incidencias registradas", len(progreso.get("incidencias", []))),
        ("Puntuación interna", "Cobertura (0-100) = peso de criterios compatibles / 100. Pesos: fecha 20, origen 20, "
                               "destino 15, vehículo 15, capacidad 15, ADR/equipamiento 5, forma de carga 5, comentarios 5. "
                               "Información (0-100) = % de 7 campos clave publicados por el camión. "
                               "Es una puntuación interna del bot, NO una probabilidad de contratación."),
        ("Advertencias", "Coincidir en provincia no demuestra proximidad. Solapar fechas no garantiza llegada a tiempo. "
                         "No se calculan kilómetros en vacío. Un mismo camión puede aparecer en varias cargas y no "
                         "está reservado. La flota declarada no acredita disponibilidad. Una oferta de camión no "
                         "demuestra que el contacto sea el conductor ni que el vehículo sea propio. "
                         "Los precios '0' se conservan literalmente y su significado no está confirmado."),
        ("Acciones realizadas en Wtransnet", "Sólo búsqueda, navegación y lectura. Ninguna publicación, contacto, "
                                             "contratación, marcado de interés ni cambio de configuración."),
    ]
    for i, (k, v) in enumerate(filas, 2):
        _texto(ws, i, 1, k).font = Font(bold=True)
        _texto(ws, i, 2, v)
    _cerrar_hoja(ws)

    ws = _hoja(wb, "Cargas", COLS_OFERTA, ANCHOS)
    _escribir_ofertas(ws, cargas)
    _cerrar_hoja(ws)

    ws = _hoja(wb, "Camiones", COLS_OFERTA, ANCHOS)
    _escribir_ofertas(ws, camiones_usados)
    _cerrar_hoja(ws)

    cols_m = ["ID carga", "ID camión", "Posición en la carga", "Estado general", "Cobertura (interna)",
              "Información (interna)", "Ruta carga", "Fechas carga", "Ruta camión", "Fechas camión",
              "Empresa carga", "Contacto carga", "Teléfonos carga", "Emails carga", "Avisos contacto carga",
              "Empresa camión", "Contacto camión", "Teléfonos camión", "Emails camión", "Fuente contacto camión",
              "Avisos contacto camión"] + [C.NOMBRES[k] for k in C.PESOS] + [
              "Información pendiente", "Motivo de la coincidencia", "Advertencias", "Ficha carga", "Ficha camión"]
    ws = _hoja(wb, "Matching", cols_m, ANCHOS)
    fila = 2
    for cid, rel in rels_ok[:100]:
        ca, cm = progreso["cargas"][cid], progreso["camiones"][rel["camion"]]
        vals = [cid, rel["camion"], rel["posicion"], rel["estado"], rel["cobertura"], rel["informacion"],
                f"{_ubic(ca['origen'])} → {_ubic(ca['destino'])}", ca["campos"].get("disponibilidad"),
                f"{_ubic(cm['origen'])} → {_ubic(cm['destino'])}", cm["campos"].get("disponibilidad"),
                ca.get("empresa_nombre"), ca.get("contacto"), ca.get("telefonos", []) + ca.get("moviles", []),
                ca.get("emails"), ca.get("avisos_contacto"), cm.get("empresa_nombre"), cm.get("contacto"),
                cm.get("telefonos", []) + cm.get("moviles", []), cm.get("emails"), _fuente(cm), cm.get("avisos_contacto")]
        vals += [rel["criterios"][k]["estado"] for k in C.PESOS]
        vals += [rel["pendiente"], rel["motivo"], rel.get("advertencias")]
        for j, v in enumerate(vals, 1):
            c = _texto(ws, fila, j, v)
            if isinstance(v, str) and v in COLOR_ESTADO:
                c.fill = PatternFill("solid", fgColor=COLOR_ESTADO[v])
        _enlace(ws, fila, len(vals) + 1, ca.get("url"))
        _enlace(ws, fila, len(vals) + 2, cm.get("url"))
        fila += 1
    ws.freeze_panes = "C2"
    _cerrar_hoja(ws)

    ws = _hoja(wb, "Incidencias", ["Momento", "Tipo", "Oferta", "Detalle", "Enlace"], ANCHOS)
    inc = list(progreso.get("incidencias", []))
    for cid in progreso["orden_cargas"]:
        for rel in progreso["relaciones"].get(cid, []):
            if rel["estado"] == "Descartado":
                inc.append({"momento": None, "tipo": "Candidato descartado", "oferta": f"{cid} ↔ {rel['camion']}",
                            "detalle": rel["descarte"], "enlace": progreso["camiones"].get(rel["camion"], {}).get("url", "")})
    for a, b, motivo in posibles_duplicados(cargas) + posibles_duplicados(list(progreso["camiones"].values())):
        inc.append({"momento": None, "tipo": "Posible duplicado (no fusionado)", "oferta": f"{a} / {b}", "detalle": motivo, "enlace": ""})
    for i, x in enumerate(inc, 2):
        _texto(ws, i, 1, x.get("momento"))
        _texto(ws, i, 2, x.get("tipo"))
        _texto(ws, i, 3, x.get("oferta"))
        _texto(ws, i, 4, x.get("detalle"))
        _enlace(ws, i, 5, x.get("enlace"), "Abrir")
    _cerrar_hoja(ws)

    wb.move_sheet("Resumen", offset=-wb.sheetnames.index("Resumen"))
    wb.save(ruta)
    return ruta


def _criterios_txt(d: dict) -> str:
    partes = []
    for k, v in (d or {}).items():
        if v in (None, "", [], {}):
            continue
        if isinstance(v, list) and v and isinstance(v[0], dict):
            v = " + ".join(", ".join(f"{a}={b}" for a, b in x.items() if b) for x in v)
        elif isinstance(v, list):
            v = ", ".join(map(str, v))
        partes.append(f"{k}: {v}")
    return "\n".join(partes) or "(ninguno)"
