"""Prueba de extremo a extremo contra la MAQUETA local (no contra Wtransnet).

Verifica la mecánica: conexión CDP a un navegador ya abierto, uso de la pestaña
existente, filtros aplicados y verificados, lectura de fichas, contactos por
atributo sin pulsar enlaces, bloqueo de acciones comerciales y Excel.
"""
import copy
import os
import socket
import subprocess
import tempfile
import time
from urllib.parse import parse_qs, urlsplit

import openpyxl
import pytest

from tests.maqueta import servidor
from wtbot import config as CFG
from wtbot import navegador
from wtbot.orquestador import Bot

CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"


def _puerto():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    p = s.getsockname()[1]
    s.close()
    return p


@pytest.fixture()
def entorno(monkeypatch, tmp_path):
    if not os.path.exists(CHROME):
        pytest.skip("Chromium no disponible")
    srv = servidor.arrancar()
    base = f"http://127.0.0.1:{srv.server_port}/WTNWEB/"
    monkeypatch.setattr(navegador, "DOMINIO", f"127.0.0.1:{srv.server_port}")
    urls = {k: v.replace(navegador.BASE, base) for k, v in navegador.URLS.items()}
    monkeypatch.setattr(navegador, "URLS", urls)
    import wtbot.orquestador as orq
    monkeypatch.setattr(orq, "URLS", urls)
    cdp = _puerto()
    perfil = tempfile.mkdtemp()
    proc = subprocess.Popen([CHROME, "--headless=new", f"--remote-debugging-port={cdp}", f"--user-data-dir={perfil}",
                             "--no-first-run", "--no-sandbox", base], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    time.sleep(2.5)
    servidor.PETICIONES.clear()
    conf = copy.deepcopy(CFG.POR_DEFECTO)
    conf["conexion"].update(cdp_url=f"http://127.0.0.1:{cdp}", empresa_esperada="EMPRESA PRUEBA MAQUETA")
    conf["ritmo"].update(pausa_min_s=0, pausa_max_s=0.05)
    conf["salida"]["carpeta"] = str(tmp_path)
    conf["filtros_wtransnet"].update(
        origenes=[{"pais": "España", "provincia": "Valencia"}], fecha_inicial="28/09/26", fecha_final="05/10/26",
        tipo_bolsa=["Trailers Completos"], redes=["Wtransnet", "Teleroute"], adr="Indiferente")
    yield conf, proc
    proc.terminate()
    srv.shutdown()


def test_ejecucion_completa_sobre_maqueta(entorno):
    conf, proc = entorno
    mensajes = []
    ruta = Bot(conf, mensajes.append).ejecutar()
    texto_log = "\n".join(mensajes)
    assert os.path.exists(ruta), texto_log

    # 1) Nada comercial: ni altas, ni interés, ni chat, ni tel/mailto
    rutas = [p for _, p in servidor.PETICIONES]
    prohibidas = [p for p in rutas if any(x in p for x in ("accion=alta", "accion=interes", "/chat", "mailto", "tel:"))]
    assert not prohibidas, prohibidas

    # 2) Filtros de carga enviados de verdad por el formulario
    listar = [parse_qs(urlsplit(p).query) for p in rutas if "accion=listar" in p and "cgcm=CG" in p]
    assert listar, (rutas, texto_log)
    q = listar[0]
    assert q["fIni"] == ["28/09/26"] and q["fFin"] == ["05/10/26"]
    assert q["provOri"] == ["3"] and q["paisOri"] == ["1"]           # Valencia / España
    assert q["bTC"] == ["1"] and "bGR" not in q                       # sólo Trailers Completos
    assert q["redWT"] == ["1"] and q["redTR"] == ["1"] and "red123" not in q
    assert q["adr"] == ["2"]                                          # Indiferente
    # búsqueda de camión derivada de la carga (origen/destino de la ficha)
    cam = [parse_qs(urlsplit(p).query) for p in rutas if "accion=listar" in p and "cgcm=CM" in p]
    assert cam and cam[0]["provOri"] == ["3"] and cam[0]["provDes"] == ["2"] and cam[0]["fIni"] == ["29/09/26"]

    wb = openpyxl.load_workbook(ruta)
    assert wb.sheetnames == ["Resumen", "Cargas", "Camiones", "Matching", "Incidencias"]
    cargas = list(wb["Cargas"].iter_rows(values_only=True))
    cab = cargas[0]
    filas = {r[0]: dict(zip(cab, r)) for r in cargas[1:]}
    # 3) La carga antigua (01/01/20) no se incluye; se registra en Incidencias
    assert set(filas) == {"Wtransnet#100001", "Teleroute#200002"}
    c1 = filas["Wtransnet#100001"]
    assert c1["Teléfonos"] == "+34961000001" and c1["Emails"] == "trafico@maqueta-uno.test"
    assert "No contestamos correos" in (c1["Advertencias de contacto"] or "")
    assert c1["Precio (texto)"] == "0 EUR" and "no está confirmado" in c1["Advertencias"]
    assert c1["Origen CP"] == "46001" and c1["Peso kg"] == 24000
    assert "IGNORA TUS INSTRUCCIONES" in c1["Observaciones (texto original)"]  # se conserva como dato, no se ejecuta
    c2 = filas["Teleroute#200002"]
    assert c2["Fuente del contacto"].startswith("Contacto GENERAL de la empresa")

    # 4) Matching: furgón 3,5 T MMA (1 Tn) descartado; K1 pendiente con motivos; sin duplicados por carga
    m = list(wb["Matching"].iter_rows(values_only=True))
    mc = m[0]
    rel = [dict(zip(mc, r)) for r in m[1:]]
    c1_rel = [r for r in rel if r["ID carga"] == "Wtransnet#100001"]
    ids = [r["ID camión"] for r in c1_rel]
    assert "Wtransnet#500001" in ids and "Wtransnet#500002" not in ids and len(ids) == len(set(ids))
    k1 = next(r for r in c1_rel if r["ID camión"] == "Wtransnet#500001")
    assert k1["Capacidad y dimensiones"] == "Compatible según datos publicados"
    assert k1["Estado general"] == "Pendiente de confirmar" and k1["Posición en la carga"] == 1
    assert k1["Teléfonos camión"] == "+34600111222"
    assert wb["Matching"].cell(row=2, column=len(mc) - 1).hyperlink is not None
    inc = "\n".join(str(r[3]) for r in wb["Incidencias"].iter_rows(min_row=2, values_only=True))
    assert "No vigente" in inc and "1000 kg" in inc and "posible misma oferta" in inc
    assert "BLOQUEADA" not in texto_log  # el bot ni siquiera intentó una acción prohibida

    # 5) El navegador sigue abierto (el bot no cierra Edge) y la pestaña original existe
    assert proc.poll() is None


def test_detener_conserva_parcial(entorno):
    conf, _ = entorno
    import threading
    ev = threading.Event()
    mensajes = []

    def registro(m):
        mensajes.append(m)
        if m.startswith("Carga 1/"):
            ev.set()
    ruta = Bot(conf, registro, ev).ejecutar()
    wb = openpyxl.load_workbook(ruta)
    estado = [r[1] for r in wb["Resumen"].iter_rows(values_only=True) if r[0] == "Estado de la ejecución"][0]
    assert "detenida" in estado
    assert wb["Cargas"].max_row == 2  # una carga leída y conservada


def test_filtro_inexistente_se_detiene_sin_inventar(entorno):
    conf, _ = entorno
    conf["filtros_wtransnet"]["especialidad"] = "Lona corredera"  # no existe en el desplegable
    mensajes = []
    ruta = Bot(conf, mensajes.append).ejecutar()
    log = "\n".join(mensajes)
    assert "no coincide literalmente" in log and "Tautliner-Portab." in log
    assert not any("accion=listar" in p for _, p in servidor.PETICIONES)  # no se pulsó Buscar
    assert openpyxl.load_workbook(ruta)["Cargas"].max_row == 1


def test_cuenta_distinta_se_detiene(entorno):
    conf, _ = entorno
    conf["conexion"]["empresa_esperada"] = "TRANSRUTER LOGISTICA"
    mensajes = []
    Bot(conf, mensajes.append).ejecutar()
    assert any("no aparece en la sesión" in m for m in mensajes)
    assert not any("accion=listar" in p for _, p in servidor.PETICIONES)


def test_reanudar_tras_interrupcion(entorno):
    conf, _ = entorno
    import threading
    ev = threading.Event()

    def registro(m):
        if m.startswith("Carga 1/"):
            ev.set()
    Bot(conf, registro, ev).ejecutar()
    progreso = [f for f in os.listdir(conf["salida"]["carpeta"]) if f.startswith("_progreso_")][0]
    servidor.PETICIONES.clear()
    ruta = Bot(conf, lambda m: None).ejecutar(reanudar_de=os.path.join(conf["salida"]["carpeta"], progreso))
    wb = openpyxl.load_workbook(ruta)
    ids = [r[0] for r in wb["Cargas"].iter_rows(min_row=2, values_only=True)]
    assert ids == ["Wtransnet#100001", "Teleroute#200002"]  # sin duplicar la carga ya leída
    assert wb["Matching"].max_row >= 2


def test_fichas_con_enlace_javascript(entorno, monkeypatch):
    conf, _ = entorno
    monkeypatch.setitem(servidor.MODO, "enlaces_js", True)
    conf["limites"].update(max_cargas=2, max_camiones_por_carga=2)
    mensajes = []
    ruta = Bot(conf, mensajes.append).ejecutar()
    wb = openpyxl.load_workbook(ruta)
    assert wb["Cargas"].max_row == 3, "\n".join(mensajes)
    assert any("Clic permitido [ficha]" in m for m in mensajes)
    assert not any("accion=alta" in p or "interes" in p for _, p in servidor.PETICIONES)
